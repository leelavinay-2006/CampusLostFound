from flask import Flask, render_template, request, redirect, url_for, flash
from openpyxl import Workbook, load_workbook
from datetime import datetime
from pathlib import Path

app = Flask(__name__)
app.secret_key = "campus-lost-found-local-key"

BASE_DIR = Path(__file__).resolve().parent
EXCEL_FILE = BASE_DIR / "lost_found.xlsx"

HEADERS = [
    "ID", "Status", "Item Name", "Description",
    "Location", "Date", "Contact", "Submitted At"
]


def setup_excel():
    if not EXCEL_FILE.exists():
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Reports"
        sheet.append(HEADERS)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = "A1:H1"
        workbook.save(EXCEL_FILE)


def get_reports():
    setup_excel()
    workbook = load_workbook(EXCEL_FILE, read_only=True, data_only=True)
    sheet = workbook["Reports"]
    reports = list(sheet.iter_rows(min_row=2, values_only=True))
    workbook.close()
    return [row for row in reports if row[0] is not None]


@app.route("/", methods=["GET", "POST"])
def home():
    if request.method == "POST":
        status = request.form.get("status", "").strip().lower()
        item_name = request.form.get("item_name", "").strip()
        description = request.form.get("description", "").strip()
        location = request.form.get("location", "").strip()
        date = request.form.get("date", "").strip()
        contact = request.form.get("contact", "").strip()

        if status not in ("lost", "found"):
            flash("Please choose Lost or Found.")
        elif not all([item_name, description, location, date]):
            flash("Please fill in all required fields.")
        else:
            setup_excel()
            workbook = load_workbook(EXCEL_FILE)
            sheet = workbook["Reports"]

            next_id = max(
                (int(row[0]) for row in sheet.iter_rows(
                    min_row=2, max_col=1, values_only=True
                ) if row[0] is not None),
                default=0
            ) + 1

            sheet.append([
                next_id, status.title(), item_name,
                description, location, date, contact,
                datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            ])

            workbook.save(EXCEL_FILE)
            workbook.close()
            flash(f"Report #{next_id} saved successfully to Excel!")
            return redirect(url_for("home"))

    reports = get_reports()
    query = request.args.get("q", "").strip().lower()
    status_filter = request.args.get("status", "").strip().lower()

    if query:
        reports = [
            row for row in reports
            if any(query in str(value or "").lower() for value in row[1:6])
        ]

    if status_filter in ("lost", "found"):
        reports = [
            row for row in reports
            if str(row[1]).lower() == status_filter
        ]

    return render_template("index.html", reports=reports)


if __name__ == "__main__":
    setup_excel()
    app.run(debug=True)